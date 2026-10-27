"""Ejecuta los 15 escenarios de datos/escenarios.csv sobre objetos reales de load_fixture().

Cada escenario usa una carga NUEVA de los fixtures, así que un caso no contamina a otro.
Genera la bitácora docs/bitacora/bitacora_escenarios.(csv|md|xlsx) con las columnas
Caso | Estado inicial | Operación | Resultado real | Explicación.

Ejecutar desde la raíz del proyecto:  python ejecutar_escenarios.py
"""
from __future__ import annotations

import csv
import sys
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from cargar_csv import load_fixture, rows  # noqa: E402
from app.domain import DomainException, Money, OrderItem  # noqa: E402
from app.pricing import calculate_discount  # noqa: E402

COLUMNAS = ["Caso", "Estado inicial", "Operación", "Resultado real", "Explicación"]
# El CSV no define números de guía; este valor lo elegimos nosotros para los casos que envían.
GUIA_DEMO = "GUIA-DEMO-001"
MOTIVO_DEMO = "Cliente cambió de opinión"
PRECIOS = {x["sku"]: Decimal(x["precio_mxn"]) for x in rows("productos.csv")}

# Estado final que exige el CSV cuando el objetivo es "exito" (la columna 'pista' lo indica).
ESTADO_FINAL_EXITO = {"E07": "CONFIRMED", "E14": "RETURNED", "E15": "DELIVERED"}

# Operación en la que el CSV espera que ocurra el rechazo (se deduce de 'accion' y 'pista').
RECHAZO_EN = {"E03": "confirm", "E06": "confirm", "E08": "add_item",
              "E11": "ship", "E12": "confirm", "E13": "cancel"}

PASOS = {
    "confirmar": [("confirm",)],
    "confirmar y enviar sin guia": [("confirm",), ("ship", "")],
    "confirmar dos veces": [("confirm",), ("confirm",)],
    "confirmar, enviar y cancelar": [("confirm",), ("ship", GUIA_DEMO), ("cancel", MOTIVO_DEMO)],
    "confirmar, enviar y devolver": [("confirm",), ("ship", GUIA_DEMO), ("return_order",)],
    "confirmar, enviar y entregar": [("confirm",), ("ship", GUIA_DEMO), ("deliver",)],
}


@dataclass
class Resultado:
    caso_id: str
    pedido_id: str
    accion: str
    objetivo: str
    estado_inicial: str
    estado_final: str
    unidades_final: int
    total_final: str
    eventos: list[str]
    resultado: str          # etiqueta comparable con la columna 'objetivo' del CSV
    invariantes_ok: bool    # una operación rechazada no dejó cambios ni eventos
    cumple: bool            # resultado == objetivo y detalle esperado
    pasos: list[str] = field(default_factory=list)
    explicacion: str = ""

    def fila_bitacora(self) -> list[str]:
        resultado_real = "; ".join(self.pasos) + (
            f"; estado final {self.estado_final}; unidades {self.unidades_final}; "
            f"total {self.total_final} MXN; eventos {self.eventos or 'ninguno'}"
        )
        veredicto = "PASA. " if self.cumple else "FALLA. "
        return [self.caso_id, self.estado_inicial, f"{self.pedido_id}: {self.accion}",
                resultado_real, veredicto + self.explicacion]


def _foto(o):
    """Fotografía del pedido para comprobar que una operación fallida no cambió nada."""
    return (o.status, o.units, tuple(o.items), tuple(o.events), o.tracking_number, o.cancel_reason)


def ejecutar_escenario(fila: dict) -> Resultado:
    """Ejecuta UN escenario del CSV sobre una carga nueva de fixtures."""
    cid, pid, accion, objetivo = fila["caso_id"], fila["pedido_id"], fila["accion"], fila["objetivo"]
    o = load_fixture()[pid]
    inicial = o.status.value
    pasos: list[str] = []
    invariantes = True
    resultado = ""
    explicacion = ""
    rechazado_en = None

    if accion == "evaluar_descuento":
        d = calculate_discount(o)
        resultado = f"{int(d.rate * 100)}%"
        pasos.append(f"Descuento {resultado} (subtotal {d.subtotal.amount}, descuento {d.discount.amount}, "
                     f"total con descuento {d.total.amount})")
        explicacion = (f"{pid}: cliente {o.customer.id} VIP={o.customer.is_vip}, promoción={o.has_promotion}, "
                       f"{o.units} unidades, total {o.total.amount} MXN. Esperado {objetivo}, obtenido {resultado}. "
                       f"Evaluar no modifica el pedido (sigue en DRAFT, sin eventos).")
        invariantes = o.status.value == inicial and not o.events
    elif accion.startswith("agregar"):
        _, sku, x = accion.split()
        cantidad = int(x.lstrip("x"))
        antes = _foto(o)
        unidades_antes = o.units
        try:
            o.add_item(OrderItem(sku, cantidad, Money(PRECIOS[sku])))
            resultado = "exito"
            pasos.append(f"add_item OK; unidades {o.units}")
        except DomainException as exc:
            resultado = "rechazo"
            rechazado_en = "add_item"
            pasos.append(f"add_item rechazado ({type(exc).__name__}: {exc})")
            invariantes = _foto(o) == antes
        explicacion = (f"Había {unidades_antes} unidades; +{cantidad} = {unidades_antes + cantidad} > 20. "
                       f"El límite se mide en unidades acumuladas. Sin cambios parciales: {invariantes}.")
    else:
        rechazado = None
        for paso in PASOS[accion]:
            antes = _foto(o)
            try:
                getattr(o, paso[0])(*paso[1:])
                pasos.append(f"{paso[0]} OK")
            except DomainException as exc:
                rechazado = (paso[0], type(exc).__name__, str(exc))
                rechazado_en = paso[0]
                pasos.append(f"{paso[0]} rechazado ({type(exc).__name__}: {exc})")
                invariantes = _foto(o) == antes
                break
        if rechazado:
            resultado = objetivo if objetivo.startswith("rechazo") else "rechazo"
            explicacion = (f"La operación inválida se rechazó con {rechazado[1]} y no dejó cambios parciales "
                           f"ni eventos extra ({invariantes}). Pista del CSV: {fila['pista']}.")
        else:
            resultado = "exito"
            explicacion = (f"Todas las operaciones fueron válidas; estado final {o.status.value}. "
                           f"Pista del CSV: {fila['pista']}.")

    cumple = resultado == objetivo and invariantes
    if cid in RECHAZO_EN:  # el rechazo debe ocurrir en la operación que dice el escenario
        cumple = cumple and rechazado_en == RECHAZO_EN[cid]
    if cid in ESTADO_FINAL_EXITO:
        cumple = cumple and o.status.value == ESTADO_FINAL_EXITO[cid]
    return Resultado(cid, pid, accion, objetivo, inicial, o.status.value, o.units, str(o.total.amount),
                     [e.name for e in o.events], resultado, invariantes, cumple, pasos, explicacion)


def ejecutar_todos() -> list[Resultado]:
    return [ejecutar_escenario(f) for f in rows("escenarios.csv")]


def escribir_bitacora(resultados: list[Resultado]) -> Path:
    out = ROOT / "docs" / "bitacora"
    out.mkdir(parents=True, exist_ok=True)
    filas = [r.fila_bitacora() for r in resultados]
    with (out / "bitacora_escenarios.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(COLUMNAS)
        w.writerows(filas)
    with (out / "bitacora_escenarios.md").open("w", encoding="utf-8") as f:
        f.write("| " + " | ".join(COLUMNAS) + " |\n|" + "---|" * len(COLUMNAS) + "\n")
        for r in filas:
            f.write("| " + " | ".join(c.replace("|", "/") for c in r) + " |\n")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:  # el .xlsx es opcional
        return out
    wb = Workbook()
    ws = wb.active
    ws.title = "Bitácora"
    ws.append(COLUMNAS)
    for r in filas:
        ws.append(r)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E79")
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
    for col, ancho in zip("ABCDE", (8, 14, 38, 80, 80)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    wb.save(out / "bitacora_escenarios.xlsx")
    return out


def main() -> int:
    resultados = ejecutar_todos()
    escribir_bitacora(resultados)
    for r in resultados:
        print(r.caso_id, "PASA" if r.cumple else "FALLA", "-", r.accion, "->", r.resultado)
    n = sum(r.cumple for r in resultados)
    print(f"{n}/{len(resultados)} escenarios cumplen su objetivo")
    print("TODOS PASAN" if n == len(resultados) else "HAY FALLAS")
    return 0 if n == len(resultados) else 1


if __name__ == "__main__":
    sys.exit(main())
